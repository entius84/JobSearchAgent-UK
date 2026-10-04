import os
import json
from typing import Dict, List
from docx import Document
import openpyxl

VAULT_DIR = r"C:\JobSearchAgent\CV certificates and CPD"


def read_txt_or_md(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        return f"[Error reading text file {file_path}: {e}]"


def read_docx(file_path: str) -> str:
    try:
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[Error reading docx {file_path}: {e}]"


def read_xlsx(file_path: str) -> str:
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text_data: List[str] = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                if any(row):
                    text_data.append(
                        " | ".join(str(c) if c is not None else "" for c in row)
                    )
        return "\n".join(text_data)
    except Exception as e:
        return f"[Error reading xlsx {file_path}: {e}]"


def load_vault_documents(vault_path: str = VAULT_DIR) -> Dict[str, str]:
    """
    Walk the local vault and return a dict:
      key   = relative path (for traceability)
      value = extracted plain text content

    This is the ONLY context fed to the model, to reduce hallucinations.
    """
    documents: Dict[str, str] = {}

    if not os.path.exists(vault_path):
        print(f"[Extractor] Vault path not found: {vault_path}")
        return documents

    for root, _, files in os.walk(vault_path):
        for filename in files:
            file_path = os.path.join(root, filename)
            rel_path = os.path.relpath(file_path, vault_path)

            # Skip temp Office files
            if filename.startswith("~$"):
                continue

            ext = os.path.splitext(filename)[1].lower()

            if ext in (".txt", ".md"):
                content = read_txt_or_md(file_path)
            elif ext == ".docx":
                content = read_docx(file_path)
            elif ext == ".xlsx":
                content = read_xlsx(file_path)
            else:
                # Ignore unknown formats to keep context clean
                continue

            documents[rel_path] = content

    return documents


def build_vault_context(docs: Dict[str, str]) -> str:
    """
    Build a compact, structured context string for the LLM.
    Each document is clearly delimited to reduce blending and hallucinations.
    """
    chunks: List[str] = []
    for name, content in docs.items():
        # Keep headers short and explicit
        chunks.append(f"--- FILE: {name} ---\n{content}\n")
    return "\n".join(chunks)


if __name__ == "__main__":
    docs = load_vault_documents()
    print(f"[Extractor] Loaded {len(docs)} documents from vault.")
    context = build_vault_context(docs)
    # Optional: write to a cache file for debugging / reuse
    cache_path = os.path.join(VAULT_DIR, "vault_context_cache.txt")
    with open(cache_path, "w", encoding="utf-8") as f:
        f.write(context)
    print(f"[Extractor] Context cache written to: {cache_path}")
