import os
from docx import Document
from pypdf import PdfReader
import requests

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "qwen2.5:14b"

VAULT_DIRS = [r"C:\JobSearchAgent\CV certificates and CPD"]
OUTPUT_FILE = os.path.join(VAULT_DIRS[0], "master_cv_and_cpd.md")


def read_docx(file_path):
  try:
    doc = Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])
  except Exception as e:
    return f"[Error reading docx {file_path}: {e}]"


def read_pdf(file_path):
  try:
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
      extracted = page.extract_text()
      if extracted:
        text += extracted + "\n"
    return text
  except Exception as e:
    return f"[Error reading pdf {file_path}: {e}]"


def load_vault_documents(vault_paths):
  documents = {}
  for vault_path in vault_paths:
    if not os.path.exists(vault_path):
      print(f"Warning: Directory not found -> {vault_path}")
      continue
    print(f"Scanning directory and subfolders: {vault_path}")
    for root, dirs, files in os.walk(vault_path):
      for filename in files:
        file_path = os.path.join(root, filename)
        rel_path = os.path.relpath(file_path, vault_path)
        key = f"{os.path.basename(vault_path)}/{rel_path}"

        if filename.endswith((".txt", ".md")):
          try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
              documents[key] = f.read()
          except Exception as e:
            print(f"Skipping {filename}: {e}")
        elif filename.endswith(".docx") and not filename.startswith("~$"):
          documents[key] = read_docx(file_path)
        elif filename.endswith(".pdf"):
          documents[key] = read_pdf(file_path)
  return documents


def query_ollama(prompt, context):
  system_instruction = (
      "CRITICAL: You are writing a comprehensive Master CV and CPD record for "
      "Vincenzo Vitiello, a mechanical engineer holding a BEng (Hons) from The "
      "Open University and AMIMechE status, with extensive NHS operational, "
      "quality management, and decontamination background. Extract and synthesize "
      "all real facts, dates, qualifications, and employment history from the "
      "provided documents. Never use placeholder brackets like [Your Name]; "
      "always use Vincenzo's actual details."
  )
  full_prompt = (
      f"{system_instruction}\n\nCareer Vault Documents:\n{context}\n\nTask:\n{prompt}"
  )
  payload = {"model": MODEL_NAME, "prompt": full_prompt, "stream": False}
  print(
      "Synthesizing Master CV and CPD with Ollama from local folder (this may"
      " take a minute)..."
  )
  response = requests.post(OLLAMA_URL, json=payload)
  if response.status_code == 200:
    return response.json().get("response", "")
  else:
    raise Exception(f"Error: {response.status_code} - {response.text}")


if __name__ == "__main__":
  print("Loading vault documents from C:\\JobSearchAgent\\CV certificates and CPD...")
  docs = load_vault_documents(VAULT_DIRS)
  print(f"Total readable documents loaded: {len(docs)}")
  vault_context = "\n\n".join(
      [f"--- {name} ---\n {content}" for name, content in docs.items()]
  )

  prompt = (
      "Generate a professional, comprehensive Master CV and Continuing "
      "Professional Development (CPD) record in Markdown format. Include:\n"
      "1. Professional Summary (highlighting mechanical engineering, BEng, "
      "AMIMechE, QMS coordination, and NHS operational/decontamination "
      "leadership).\n"
      "2. Core Competencies & Technical Skills.\n"
      "3. Professional Experience (detailed chronological breakdown of roles "
      "at Buckinghamshire Healthcare NHS Trust and prior experience).\n"
      "4. Education & Certifications (Open University BEng Hons, ILM Level 5 "
      "Leadership & Management, Train the Healthcare Trainer, etc.).\n"
      "5. CPD Record & Technical Projects (incorporating all professional "
      "development activities from the folder).\n"
      "Ensure all data is pulled directly from the provided vault context with "
      "no placeholders."
  )

  master_content = query_ollama(prompt, vault_context)

  with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
    f.write(master_content)

  print(f"Successfully generated and saved: {OUTPUT_FILE}")