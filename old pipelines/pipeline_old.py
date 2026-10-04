import json
import os
import time
from docx import Document
import openpyxl
from playwright.sync_api import sync_playwright
import requests

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "qwen2.5:14b"
PROCESSED_FILE = "processed_jobs.json"
APPLICATIONS_DIR = "tailored_applications"


def load_processed_jobs():
  if os.path.exists(PROCESSED_FILE):
    try:
      with open(PROCESSED_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
        # Automatically handle legacy list format if it exists
        if isinstance(data, list):
          return {url: "N/A" for url in data}
        return data
    except Exception:
      return {}
  return {}


def save_processed_job(url, score):
  processed = load_processed_jobs()
  processed[url] = score
  with open(PROCESSED_FILE, "w", encoding="utf-8") as f:
    json.dump(processed, f, indent=2)


def query_ollama(prompt, context=""):
  full_prompt = f"Context (My Career Vault):\n{context}\n\nQuery:\n{prompt}"
  payload = {"model": MODEL_NAME, "prompt": full_prompt, "stream": False}
  response = requests.post(OLLAMA_URL, json=payload)
  if response.status_code == 200:
    return response.json().get("response", "")
  else:
    raise Exception(f"Error: {response.status_code} - {response.text}")


def read_docx(file_path):
  try:
    doc = Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])
  except Exception as e:
    return f"[Error reading docx: {e}]"


def read_gdoc(file_path):
  try:
    with open(file_path, "r", encoding="utf-8") as f:
      data = json.load(f)
      return f"[Google Doc Link: {data.get('url', 'N/A')}]"
  except Exception as e:
    return f"[Error reading gdoc: {e}]"


def read_xlsx(file_path):
  try:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text_data = []
    for sheet in wb.sheetnames:
      ws = wb[sheet]
      for row in ws.iter_rows(values_only=True):
        if any(row):
          text_data.append(
              " | ".join([str(c) if c is not None else "" for c in row])
          )
    return "\n".join(text_data)
  except Exception as e:
    return f"[Error reading xlsx: {e}]"


def load_vault_documents(vault_path):
  documents = {}
  if not os.path.exists(vault_path):
    return documents
  for root, dirs, files in os.walk(vault_path):
    for filename in files:
      file_path = os.path.join(root, filename)
      rel_path = os.path.relpath(file_path, vault_path)
      if filename.endswith((".txt", ".md")):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
          documents[rel_path] = f.read()
      elif filename.endswith(".docx") and not filename.startswith("~$"):
        documents[rel_path] = read_docx(file_path)
      elif filename.endswith(".gdoc"):
        documents[rel_path] = read_gdoc(file_path)
      elif filename.endswith(".xlsx"):
        documents[rel_path] = read_xlsx(file_path)
  return documents


def save_job_on_linkedin(page):
  try:
    save_btn = page.locator(
        "button.jobs-save-button, button:has-text('Save')"
    ).first
    if save_btn.count() > 0:
      aria_label = save_btn.get_attribute("aria-label") or ""
      if "saved" not in aria_label.lower():
        save_btn.click()
        print("-> Added to LinkedIn Saved jobs tracker.")
        time.sleep(1)
      else:
        print("-> Job is already saved on LinkedIn.")
  except Exception as e:
    print(f"-> Could not click LinkedIn Save button: {e}")


def scrape_top_applicant_jobs(max_pages=2):
  profile_dir = r"C:\JobSearchAgent\playwright_profile"
  processed_urls = load_processed_jobs()
  detailed_jobs = []
  seen_urls = set()

  p = sync_playwright().start()
  context = p.chromium.launch_persistent_context(
      user_data_dir=profile_dir, headless=False, slow_mo=50
  )
  page = context.new_page()

  for page_num in range(max_pages):
    base_url = "https://www.linkedin.com/jobs/collections/recommended/?recommendationChannel=TOP_APPLICANT"
    current_url = (
        f"{base_url}&start={page_num * 25}" if page_num > 0 else base_url
    )

    print(f"Navigating to Top Applicant collection (Page {page_num + 1})...")
    page.goto(current_url, wait_until="domcontentloaded", timeout=60000)

    if "login" in page.url or "uas" in page.url or "microsoft" in page.url:
      print("Please complete login in the browser window...")
      input("Press Enter here in the terminal *after* you are fully logged in...")

    time.sleep(3)
    print("Scrolling the inner job list container...")
    for _ in range(8):
      try:
        page.locator(
            ".jobs-search-results-list, .scaffold-layout__list"
        ).evaluate("node => node.scrollBy(0, 1500)")
      except Exception:
        page.mouse.wheel(0, 2000)
      time.sleep(2)

    cards = page.locator(".job-card-container, li.occludable-update").all()
    print(
        f"Found {len(cards)} listings on page {page_num + 1}. Filtering"
        " against history..."
    )

    page_new_found = 0
    for i, card in enumerate(cards):
      try:
        link_elem = card.locator("a[href*='/jobs/view/']").first
        if link_elem.count() == 0:
          continue
        href = link_elem.get_attribute("href")
        clean_url = href.split("?")[0] if href else ""
        full_url = (
            clean_url
            if clean_url.startswith("http")
            else f"https://www.linkedin.com{clean_url}"
        )

        if full_url in processed_urls or full_url in seen_urls:
          continue

        seen_urls.add(full_url)
        card.scroll_into_view_if_needed()
        card.click()
        time.sleep(2.5)

        title_elem = card.locator(
            ".job-card-list__title, .job-card-container__link"
        ).first
        title = (
            title_elem.inner_text().strip()
            if title_elem.count() > 0
            else "Unknown Title"
        )
        if "\n" in title:
          title = title.split("\n")[0].strip()

        date_posted = "N/A"
        time_elem = card.locator("time, .job-card-container__listed-time").first
        if time_elem.count() > 0:
          date_posted = time_elem.inner_text().strip()

        desc_elem = page.locator(
            ".jobs-description__content, .jobs-box__html-content"
        ).first
        description = (
            desc_elem.inner_text().strip()
            if desc_elem.count() > 0
            else "Description not loaded."
        )

        detailed_jobs.append({
            "title": title,
            "date": date_posted,
            "url": full_url,
            "description": description,
            "card": card,
        })
        page_new_found += 1
        print(f"[{len(detailed_jobs)}] Loaded new job for evaluation: {title}")

        if len(detailed_jobs) >= 10:
          break
      except Exception as e:
        print(f"Error scraping card {i+1}: {e}")

      if len(detailed_jobs) >= 10:
        break

    if len(detailed_jobs) >= 10 or page_new_found == 0:
      break

  return detailed_jobs, page, context, p


if __name__ == "__main__":
  print("Loading your career vault...")
  vault_dir = r"G:\Il mio Drive\lavoro"
  docs = load_vault_documents(vault_dir)
  vault_context = "\n\n".join(
      [f"--- {name} ---\n {content}" for name, content in docs.items()]
  )
  print(f"Loaded {len(docs)} files from your career vault.")

  print("Fetching new Top Applicant listings...")
  new_jobs, page, context, p_instance = scrape_top_applicant_jobs()
  print(f"Found {len(new_jobs)} new unevaluated job listings.")

  os.makedirs(APPLICATIONS_DIR, exist_ok=True)
  summary_records = []

  for i, job in enumerate(new_jobs):
    print(
        f"\nEvaluating Job {i+1}/{len(new_jobs)}: {job['title']}\nURL:"
        f" {job['url']}"
    )

    eval_prompt = (
        "Review this job description against my career vault. Provide a"
        " compatibility score out of 10, key strengths, potential gaps, and a"
        f" recommendation on whether to apply:\n\nTitle:"
        f" {job['title']}\nDescription:\n{job['description']}"
    )
    analysis = query_ollama(eval_prompt, vault_context)
    print(analysis)

    is_top_match = False
    score_str = "N/A"
    for line in analysis.split("\n"):
      if "score" in line.lower():
        score_str = line.strip()
        if any(str(num) in line for num in ["8/", "9/", "10/"]):
          is_top_match = True
        break

    summary_records.append({
        "title": job["title"],
        "score": score_str,
        "match": "🔥 High Match" if is_top_match else "Skipped / Low",
        "url": job["url"],
    })

    if is_top_match:
      print("🔥 High compatibility match detected!")
      try:
        job["card"].scroll_into_view_if_needed()
        job["card"].click()
        time.sleep(1.5)
        save_job_on_linkedin(page)
      except Exception as e:
        print(f"Could not auto-save on LinkedIn: {e}")

      print("Generating tailored CV positioning & cover letter draft...")
      cv_prompt = (
          "Based on my career vault, write tailored CV highlight bullet points"
          " and a brief cover letter pitch customized specifically for this"
          f" role:\n\nJob Title: {job['title']}\nDescription:\n{job['description']}"
      )
      tailored_content = query_ollama(cv_prompt, vault_context)

      safe_title = "".join(
          c if c.isalnum() else "_" for c in job["title"]
      ).strip("_")
      file_path = os.path.join(
          APPLICATIONS_DIR, f"application_{i+1}_{safe_title}.md"
      )
      with open(file_path, "w", encoding="utf-8") as f:
        f.write(f"# Role: {job['title']}\n")
        f.write(f"- **URL:** {job['url']}\n")
        f.write(f"- **Date Posted:** {job['date']}\n\n")
        f.write(f"## Evaluation & Fit\n{analysis}\n\n")
        f.write(f"## Tailored CV & Cover Letter Package\n{tailored_content}\n")
      print(f"Saved tailored application package to {file_path}")

    print("-" * 60)
    save_processed_job(job["url"], score_str)

  summary_path = os.path.join(APPLICATIONS_DIR, "master_summary.md")
  with open(summary_path, "w", encoding="utf-8") as f:
    f.write("# 📊 Job Application Master Dashboard\n\n")
    f.write("| Job Title | Compatibility & Score | Status | Direct Link |\n")
    f.write("| :--- | :--- | :--- | :--- |\n")
    for rec in summary_records:
      f.write(
          f"| {rec['title']} | {rec['score']} | {rec['match']} |"
          f" [View Job]({rec['url']}) |\n"
      )
  print(f"\nMaster summary dashboard updated: {summary_path}")

  try:
    context.close()
    p_instance.stop()
  except Exception:
    pass
  print("Batch processing complete!")